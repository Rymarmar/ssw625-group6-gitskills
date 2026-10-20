"""Build the RQ1 manual-analysis coding sheet (Excel) from the GitSkills sample.

Usage: python make_coding_sheet.py agent_skills_sample.db rq1_coding_sheet.xlsx
Sample (seed 625): 20 most widely reused contents, 20 random other cross-repo
contents, 10 random single-repository contents.
"""
import sqlite3
import sys
from urllib.parse import quote

import pandas as pd

db, out = sys.argv[1], sys.argv[2]
con = sqlite3.connect(db)
occ = pd.read_sql("SELECT repo_full_name, path, location_class, file_sha FROM artifacts", con)
rep = pd.read_sql("""SELECT file_sha, name, description, content, body_chars, has_scripts, has_references
                     FROM artifacts WHERE dedup_primary = 1""", con)
g = occ.groupby("file_sha").agg(n_repos=("repo_full_name", "nunique"), copies=("path", "size")).reset_index()
g = g.merge(rep, on="file_sha", how="left").sort_values(["n_repos", "copies", "file_sha"], ascending=[False, False, True])

top = g.head(20).assign(group="A: top-20 reused")
rest = g.iloc[20:]
cross = rest[rest.n_repos > 1].sample(20, random_state=625).assign(group="B: random cross-repo")
single = rest[rest.n_repos == 1].sample(10, random_state=625).assign(group="C: random single-repo")
s = pd.concat([top, cross, single]).reset_index(drop=True)


def examples(sha):
    rows = occ[occ.file_sha == sha].sort_values(["repo_full_name", "path"]).head(5)
    paths = "\n".join(f"{r.repo_full_name} : {r.path}" for r in rows.itertuples())
    links = [f"https://github.com/{r.repo_full_name}/blob/HEAD/{quote(r.path)}" for r in rows.itertuples()]
    return paths, links


ex = s.file_sha.map(examples)
sheet = pd.DataFrame({
    "id": range(1, len(s) + 1),
    "group": s.group,
    "name": s.name,
    "description": s.description,
    "n_repos": s.n_repos,
    "copies": s.copies,
    "example_locations (repo : path, up to 5)": [e[0] for e in ex],
    "link_1": [e[1][0] if len(e[1]) > 0 else "" for e in ex],
    "link_2": [e[1][1] if len(e[1]) > 1 else "" for e in ex],
    "content_preview (first 800 chars)": s.content.fillna("").str.slice(0, 800),
    "Q1 skill_type": "",
    "Q2 task_category": "",
    "Q3 reuse_channel": "",
    "Q4 notes": "",
    "file_sha": s.file_sha,
})

codebook = pd.DataFrame([
    ("Q1 skill_type", "general", "通用：任何项目都能用（如 React 最佳实践、写 commit message）"),
    ("Q1 skill_type", "project-specific", "项目专用：提到某个具体项目的文件、命令、业务"),
    ("Q1 skill_type", "unclear", "看不出来 / 内容太短 / 只是链接"),
    ("Q2 task_category", "frontend / backend / devops / testing / docs / data-ml / agent-meta / domain / other",
     "任务类别，选最贴切的一个；agent-meta = 关于 agent 本身（写 skill、写规则）；domain = 非软件领域（如能源采购）"),
    ("Q3 reuse_channel", "multi-agent install", "同一内容出现在 .agents/ .claude/ .codex/ .trae/ 等多个 agent 配置目录"),
    ("Q3 reuse_channel", "collection/mirror", "出现在 awesome-xxx、mirror、skills 合集类仓库或路径里"),
    ("Q3 reuse_channel", "template/scaffold", "出现在 template、starter、boilerplate 类路径里"),
    ("Q3 reuse_channel", "independent adoption", "被很多普通项目各自复制进来，看不出上面三种模式"),
    ("Q3 reuse_channel", "n/a", "C 组（只在 1 个仓库）不用填"),
    ("Q4 notes", "free text", "任何有意思的发现，写进报告时可引用"),
], columns=["question", "code", "meaning"])

with pd.ExcelWriter(out, engine="openpyxl") as w:
    sheet.to_excel(w, sheet_name="coding", index=False)
    codebook.to_excel(w, sheet_name="codebook", index=False)
    ws = w.sheets["coding"]
    from openpyxl.styles import Alignment, Font, PatternFill
    widths = {"A": 5, "B": 20, "C": 26, "D": 40, "E": 8, "F": 8, "G": 60, "H": 18, "I": 18,
              "J": 70, "K": 16, "L": 16, "M": 20, "N": 30, "O": 12}
    for col, wdt in widths.items():
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (row[7], row[8]):
            if c.value:
                c.hyperlink, c.value, c.font = c.value, "open on GitHub", Font(color="0563C1", underline="single")
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in "KLMN":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "D2"
    cb = w.sheets["codebook"]
    cb.column_dimensions["A"].width, cb.column_dimensions["B"].width, cb.column_dimensions["C"].width = 18, 40, 80
print(f"Wrote {len(sheet)} rows to {out}")
